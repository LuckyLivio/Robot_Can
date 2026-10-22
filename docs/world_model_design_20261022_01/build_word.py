"""Build a compact chapter-style design, reusing the verified Office Math renderer."""
import argparse
from datetime import datetime,timezone
import hashlib
import importlib.util
import json
from pathlib import Path
import re

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt,RGBColor

HERE=Path(__file__).resolve().parent
ROOT=HERE.parent.parent
OLD=ROOT/'docs/detailed_design_20261022_01'
TITLE='基于世界模型的机器人抓取搬运系统设计与实现'
OUTPUT=HERE/(TITLE+'详细设计.docx')
spec=importlib.util.spec_from_file_location('verified_design_renderer',OLD/'build_design.py')
lib=importlib.util.module_from_spec(spec);spec.loader.exec_module(lib)
lib.HERE=HERE
lib.TITLE=TITLE


def sha(path):return hashlib.sha256(Path(path).read_bytes()).hexdigest()


class ChapterBuilder(lib.Builder):
    def cover(self):
        # A chapter document has a compact title, no separate cover or TOC.
        lib.set_fonts(self.doc.styles['Title'],18,east='黑体',bold=True)
        lib.set_fonts(self.doc.styles['Heading 1'],16,east='黑体',bold=True)
        self.doc.styles['Heading 1'].paragraph_format.page_break_before=False
        self.doc.styles['Normal'].paragraph_format.space_after=Pt(3)
        section=self.doc.sections[0]
        p=section.header.paragraphs[0];p.text=TITLE;p.alignment=WD_ALIGN_PARAGRAPH.RIGHT;p.paragraph_format.first_line_indent=Pt(0)
        for r in p.runs:r.font.size=Pt(9);r.font.color.rgb=RGBColor(0,0,0)
        p=section.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
        lib.field(p,'PAGE')
        p=self.doc.add_paragraph(TITLE,'Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(4)
        p=self.doc.add_paragraph('详细设计');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(10)
        settings=self.doc.settings.element
        update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)


def main():
    parser=argparse.ArgumentParser();parser.add_argument('--replace-generated',action='store_true');args=parser.parse_args()
    if OUTPUT.exists() and not args.replace_generated:raise FileExistsError(OUTPUT)
    text=(HERE/'详细设计.md').read_text(encoding='utf-8')
    images=[HERE/p for p in re.findall(r'!\[[^\]]*\]\(([^)]+)\)',text)]
    if any(not p.is_file() for p in images):raise FileNotFoundError('All design diagrams must exist first')
    builder=ChapterBuilder();doc=builder.render(text)
    # Keep the short state table together instead of leaving only one row.
    for row in doc.tables[0].rows[:-1]:
        for cell in row.cells:
            for p in cell.paragraphs:p.paragraph_format.keep_with_next=True
    for p in doc.paragraphs:
        if p.text.startswith('本章完成基于世界模型'):
            p.paragraph_format.keep_together=True
    doc.core_properties.title=TITLE+'详细设计'
    doc.core_properties.subject='第4章 系统设计'
    doc.save(OUTPUT)
    source_paths=[HERE/'详细设计.md',Path(__file__),OLD/'build_design.py',
                  OLD/'research/references_verified.json',ROOT/'common.py',ROOT/'world_model.py',
                  ROOT/'robot_can_v2/observations.py',ROOT/'robot_can_v2/evaluate.py',ROOT/'robot_can_v2/train_world.py',
                  ROOT/'robot_can_v2/train_policy.py',ROOT/'train_policy.py',ROOT/'collect_rollout_branches.py',
                  ROOT/'collect_branches.py',ROOT/'reproduction/multiseed_20261012_01/train_suite.py',
                  *[ROOT/f'models/multiseed_20261012_01_full_s{seed}/config.json' for seed in (17,23,41)],
                  ROOT/'reproduction/recovery_v5_20261012_01/recovery_controller.py']
    manifest={'title':TITLE,'built_utc':datetime.now(timezone.utc).isoformat(),
              'scope':'Standalone Chapter 4 detailed design; no abstract, introduction, full evaluation chapter or appendices',
              'reference_pdf_sha256':sha(Path('C:/Users/Livio/Desktop/211380127.pdf')),
              'reference_scope':'PDF pages 22-35, printed 15-28; architecture, functions, data and interfaces only',
              'original_52_page_version_preserved':True,'cjk_characters':len(re.findall(r'[\u4e00-\u9fff]',text)),
              'numbered_equations':sum(builder.math_count.values()),'office_math_objects':len(builder.math_checks),
              'tables':len(doc.tables),'figures':len(doc.inline_shapes),'references':6,
              'sources_sha256':{p.relative_to(ROOT).as_posix():sha(p) for p in source_paths},
              'image_sha256':{p.relative_to(HERE).as_posix():sha(p) for p in images},
              'pre_word_outputs_sha256':{p.name:sha(p) for p in (OUTPUT,HERE/'详细设计.md')},
              'math_checks':builder.math_checks,
              'verification_note':'Native Word export and per-page visual QA follow authoring; not included in this build receipt.'}
    (HERE/'build_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({key:manifest[key] for key in ('cjk_characters','numbered_equations','office_math_objects','tables','figures','references')},ensure_ascii=False))


if __name__=='__main__':main()
